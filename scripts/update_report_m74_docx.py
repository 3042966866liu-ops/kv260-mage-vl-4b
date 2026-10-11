from __future__ import annotations

import shutil
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


SOURCE = Path(r"C:\Users\30429\Desktop\报告 - 副本.docx")
PROJECT_ROOT = Path(r"C:\Users\30429\Desktop\TeLLMe_FPGA_2026")
OUTPUT = PROJECT_ROOT / "M74方案修订版_结题报告.docx"


def sanitize_package(source: Path, target: Path) -> None:
    """Remove the source document's dangling image relationship to ../NULL."""
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                root = etree.fromstring(data)
                for rel in list(root):
                    if rel.get("Target", "").upper().endswith("NULL"):
                        root.remove(rel)
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            zout.writestr(item, data)


def set_run_font(run, *, heading: bool = False) -> None:
    name = "黑体" if heading else "宋体"
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:eastAsia"), name)


def replace_paragraph(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, heading=paragraph.style.name.startswith("Heading"))


def find_first(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise RuntimeError(f"paragraph not found: {prefix}")


def insert_before(anchor, text: str, style: str = "Body Text"):
    paragraph = anchor.insert_paragraph_before(text, style=style)
    for run in paragraph.runs:
        set_run_font(run, heading=style.startswith("Heading"))
    return paragraph


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_m74_chapter(doc: Document) -> None:
    ref = find_first(doc, "参考文献")

    insert_before(ref, "8    基于 M74 的 FPGA 多模态模型推理加速方案", "Heading 1")
    insert_before(ref, "8.1   设计目标与软硬件协同路线", "Heading 2")
    insert_before(
        ref,
        "为增强课程智能体在本地视频理解场景中的推理能力，本项目以 AMD Kria KV260 为目标平台，"
        "面向 Microsoft Mage-VL 4B 的语言侧 Prefill 矩阵计算设计 M74 候选方案。系统采用 PS–PL "
        "协同架构：PS 负责视频预处理、模型调度、Attention 与控制，PL 负责高占比的 W4A8-G64 "
        "矩阵乘加。为保持现有稳定系统可回滚，M74 作为独立离线候选，不覆盖当前稳定启动文件。",
    )
    insert_before(
        ref,
        "M74 以六 token 并行、六权重分片和两批反量化为基本并行结构。在 M73 混合精确乘法的基础上，"
        "增加权重尺度预解码与精确移位加法路径，使可由二进制移位表达的乘法在进入主计算阵列前完成"
        "结构化展开，目标是缩短反量化关键路径，并为后续完整 Prefill 累加和 Attention 融合创造时序空间。",
    )

    insert_before(ref, "8.2   数值正确性与 HLS 仿真", "Heading 2")
    insert_before(
        ref,
        "M74 已在 Vitis HLS 2026.1 中完成 C 级仿真。测试覆盖 W2、W3、W4 三种权重量化位宽、"
        "6 个 token、6 个 shard 与 2 个反量化 batch，仿真日志返回 CSim PASS 且无错误。"
        "该结果证明候选算术路径在当前测试范围内保持了既定的舍入与输出语义，但它属于软件/RTL 生成前"
        "的功能验证，不能替代完整比特流、DMA、板端内存与端到端模型正确性测试。",
    )

    insert_before(ref, "8.3   性能与资源评估", "Heading 2")
    insert_before(
        ref,
        "HLS 报告的目标周期为 3.70 ns，估算周期为 3.239 ns，对应约 308.7 MHz 的核级估算频率。"
        "资源估算为 48 个 BRAM18K、1,248 个 DSP、172,016 个 FF 和 212,020 个 LUT。由于 HLS "
        "资源模型对该类浮点/定点混合结构可能偏保守，最终资源占用仍应以 Vivado OOC 综合与完整系统"
        "实现报告为准；当前尤其需要关注 DSP 已达到器件上限以及 LUT 估算超出器件标称容量的问题。",
    )
    insert_before(
        ref,
        "在锁定的 849-token 输入上，以 270 MHz 和零传输、零控制、零 Attention 额外开销进行理论"
        "下界估算，M74/M73 同级六 token 算术阵列的矩阵计算时间约为 7,640.645 ms；匹配的 CPU "
        "TTFT 基线为 10,481.116 ms。由此得到理论加速比约 1.372×，相对时间缩短约 27.1%。"
        "该数值只说明矩阵核存在超过 CPU 25% 的理论空间，不是 KV260 实板端到端测量；完整流水仍需"
        "计入视频预处理、Attention、数据搬运、PS–PL 同步和运行时调度。",
    )

    insert_before(ref, "8.4   阶段结论与上板验证计划", "Heading 2")
    insert_before(
        ref,
        "现阶段可确认的结论是：M74 的功能仿真通过，核级时钟估算具有余量，且理论矩阵计算下界具备"
        "超过 CPU 25% 的潜力；尚不能确认其完整资源可放、270 MHz 时序收敛、实板稳定运行或端到端"
        "性能达标。下一阶段将依次完成 Vivado OOC 综合、完整 Prefill 累加与 Attention 集成、"
        "full-overlay 布线、板端正确性门禁以及与同输入同精度 CPU 基线的重复配对测试。只有上述门禁"
        "全部通过后，才把结果升级为“实板验证结论”。",
    )

    insert_before(ref, "表8-1  M74 当前验证状态与边界", "Normal")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "验证项"
    hdr[1].text = "当前结果"
    hdr[2].text = "证据性质"
    hdr[3].text = "能否作为实板结论"
    set_repeat_table_header(table.rows[0])
    rows = [
        ("W2/W3/W4、6 token、6 shard、2 batch", "CSim PASS", "Vitis HLS 功能仿真", "否"),
        ("核级时钟", "3.239 ns（约 308.7 MHz）", "HLS 估算", "否"),
        ("资源", "48 BRAM18K / 1248 DSP / 172016 FF / 212020 LUT", "HLS 估算", "否"),
        ("相对 CPU 性能", "理论 1.372×，时间缩短约 27.1%", "零额外开销计算下界", "否"),
        ("板端端到端性能", "待测", "需完整比特流与配对门禁", "尚无结论"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_run_font(run)
    ref._p.addprevious(table._tbl)


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    with tempfile.TemporaryDirectory(prefix="m74_report_") as tmp_dir:
        sanitized = Path(tmp_dir) / "sanitized.docx"
        sanitize_package(SOURCE, sanitized)
        doc = Document(sanitized)

        replace_paragraph(
            find_first(doc, "项目名称："),
            "项目名称：面向课程智能体的多模态模型与 FPGA 边缘推理加速研究",
        )
        replace_paragraph(
            find_first(doc, "报告题目："),
            "报告题目：面向电气课程智能体的多模态理解平台与 M74 FPGA 推理加速方案",
        )
        replace_paragraph(
            find_first(doc, "中文摘要："),
            "中文摘要：本项目面向电气专业课程教学与本地多模态交互场景，构建了覆盖专业语料、"
            "电路图识别、结构化解析、自动仿真和网页交互的课程智能体平台。在算法侧，系统采用 "
            "YOLO11、传统计算机视觉与 OCR 完成元器件、导线、文字及连接关系提取，并将结构化结果"
            "接入 LTspice 自动仿真链路。在模型部署侧，项目进一步以 AMD Kria KV260 为目标平台，"
            "围绕 Microsoft Mage-VL 4B 设计 M74 六 token/六分片 W4A8-G64 Prefill 算术候选，"
            "通过权重尺度预解码与精确移位加法降低反量化关键路径。Vitis HLS 2026.1 功能仿真覆盖 "
            "W2/W3/W4、6 token、6 shard 和 2 batch，结果全部通过；核级估算周期为 3.239 ns。"
            "按 270 MHz 的零额外开销计算下界，矩阵阶段理论加速比约为 CPU 的 1.372 倍，具备超过 "
            "25% 的潜力。该性能属于仿真与理论预测，完整资源、布线、板端正确性和端到端性能仍需后续"
            "实板门禁确认。本项目形成了从专业内容、视觉理解到软硬件协同推理的完整技术路线。",
        )
        replace_paragraph(
            find_first(doc, "英文摘要:"),
            "Abstract: This project develops a multimodal course assistant for electrical-engineering "
            "education, integrating domain corpora, circuit-diagram recognition, structured parsing, "
            "automatic LTspice simulation, and web interaction. For edge inference, an M74 FPGA candidate "
            "is designed for Microsoft Mage-VL 4B on the AMD Kria KV260. The six-token, six-shard W4A8-G64 "
            "Prefill arithmetic path uses weight-scale predecoding and exact shift-add operations. Vitis HLS "
            "2026.1 C simulation passes W2/W3/W4 cases across six tokens, six shards, and two dequantization "
            "batches. The estimated kernel period is 3.239 ns. Under a zero-overhead analytical model at "
            "270 MHz, the matrix-compute floor is 7,640.645 ms versus a matched CPU TTFT of 10,481.116 ms, "
            "corresponding to a theoretical 1.372x speedup. These are simulation and analytical projections, "
            "not board-level end-to-end measurements; full implementation and matched board validation remain "
            "future gates.",
        )
        replace_paragraph(
            find_first(doc, "关键词："),
            "关键词：电气课程智能体；多模态模型；Mage-VL；FPGA；W4A8；PS–PL 协同",
        )

        objective = find_first(doc, "本项目的总体目标")
        replace_paragraph(
            objective,
            "本项目的总体目标是围绕电路课程教学与本地多模态交互场景，构建具备知识增强问答、"
            "电路图识别、网表生成、自动仿真和边缘推理加速能力的课程智能体。除语料、视觉算法和"
            "平台建设外，新增目标是在 Kria KV260 上研究 Mage-VL 4B 的 PS–PL 协同推理，形成可"
            "量化、可综合、可回滚、可与 CPU 公平比较的 FPGA 加速方案。具体研究目标包括：",
        )
        insert_before(
            find_first(doc, "围绕上述目标"),
            "5) 设计 M74 六 token/六分片 Prefill 算术核，完成数值仿真、资源与时序评估，并建立"
            "板端正确性及性能门禁。",
        )

        arch = find_first(doc, "根据项目目标，系统整体上可划分")
        replace_paragraph(
            arch,
            "根据项目目标，系统整体划分为知识层、算法层、软硬件协同推理层和应用层。知识层组织"
            "教材、习题、实验和问答资源；算法层完成电路图目标检测、结构解析、OCR 与网表生成；"
            "软硬件协同层由 ARM64 PS 承担视频预处理、模型控制与非矩阵算子，PL 承担 W4A8-G64 "
            "矩阵计算；应用层通过统一网页提供问答、图像理解、仿真和性能状态展示。",
        )

        innovation_anchor = find_first(doc, "6.4")
        insert_before(innovation_anchor, "6.3.4  面向多模态模型的 M74 预解码移位加法架构", "Heading 3")
        insert_before(
            innovation_anchor,
            "现有嵌入式多模态推理常受矩阵计算量、DDR 访存与反量化开销共同限制。M74 针对六 token "
            "Prefill 的反量化关键路径，采用权重尺度预解码和精确移位加法，将部分可结构化乘法从通用"
            "浮点路径中分离，同时保持 W2/W3/W4 功能仿真一致。该设计的创新价值在于将量化格式、"
            "并行分片、反量化和 FPGA 时序约束联合设计，而非仅把现有模型原样移植到开发板。",
        )
        insert_before(
            find_first(doc, "6.3   项目创新点分析"),
            "8) 完成 M74 候选的 W2/W3/W4 功能仿真与核级性能评估，形成可复现的 FPGA 多模态"
            "Prefill 优化路线和后续实板门禁。",
        )

        limitations = find_first(doc, "这些问题说明")
        insert_before(
            limitations,
            "此外，M74 当前仍停留在仿真与综合评估阶段。HLS 资源估算显示 DSP 达到器件上限，LUT "
            "估算也超出标称容量；理论 1.372× 加速未计入 Attention、数据搬运和运行时同步。"
            "因此本报告不把该预测表述为实板实测，后续必须以完整布线和端到端配对测试校验。",
        )

        future = find_first(doc, "推动成果转化")
        insert_before(
            future,
            "5) 完成 M74 的 Vivado OOC 与 full-overlay 实现，并在 KV260 上进行正确性、稳定性、"
            "TTFT、吞吐率和 CPU/FPGA 同精度配对测试；若未达到 1.25× 门禁，则继续优化 DSP 复用、"
            "Attention 融合和 DDR 数据流。",
        )

        add_m74_chapter(doc)
        doc.save(OUTPUT)

    print(OUTPUT)


if __name__ == "__main__":
    main()
