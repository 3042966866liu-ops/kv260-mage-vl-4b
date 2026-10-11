from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


SOURCE = Path(r"C:\Users\30429\Desktop\报告 - 副本.docx")
PROJECT_ROOT = Path(r"C:\Users\30429\Desktop\TeLLMe_FPGA_2026")
OUTPUT = PROJECT_ROOT / "M74项目重写版_结题报告.docx"


def sanitize_package(source: Path, target: Path) -> None:
    with zipfile.ZipFile(source, "r") as zin, zipfile.ZipFile(
        target, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                root = etree.fromstring(data)
                for rel in list(root):
                    if rel.get("Target", "").upper().endswith("NULL"):
                        root.remove(rel)
                data = etree.tostring(
                    root, xml_declaration=True, encoding="UTF-8", standalone=True
                )
            zout.writestr(item, data)


def font_run(run, heading: bool = False) -> None:
    name = "黑体" if heading else "宋体"
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), name)


def replace(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    font_run(run, paragraph.style.name.startswith("Heading"))


def add_p(doc: Document, text: str, style: str = "Body Text"):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    font_run(r, style.startswith("Heading"))
    return p


def add_heading(doc: Document, text: str, level: int):
    return add_p(doc, text, f"Heading {level}")


def set_repeat_header(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            border = OxmlElement(f"w:{edge}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "000000")
            borders.append(border)
        tbl_pr.append(borders)
    set_repeat_header(table.rows[0])
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    font_run(run, heading=row_index == 0)
                    run.font.size = Pt(9)
                    if row_index == 0:
                        run.bold = True
    return table


def keep_front_matter_only(doc: Document) -> None:
    # Paragraphs 0–28 contain the cover and title/abstract page. Paragraph 29
    # is a second consecutive section break in the source and creates a blank
    # page, so it is intentionally removed with the old body.
    keep = {p._p for p in doc.paragraphs[:29]}
    body = doc._element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        if child not in keep:
            body.remove(child)


def rewrite_front_matter(doc: Document) -> None:
    replace(
        doc.paragraphs[11],
        "项目名称：基于 PS–PL 协同架构的多模态大模型端侧部署与推理优化",
    )
    replace(
        doc.paragraphs[20],
        "报告题目：面向 Mage-VL 4B 的 M74 FPGA 推理加速方案设计与验证",
    )
    replace(
        doc.paragraphs[23],
        "中文摘要：本项目面向资源受限嵌入式平台上的多模态大模型部署需求，围绕“模型参数规模大、"
        "首Token计算量高、DDR访存压力大以及软硬件协同开销难以隐藏”等问题，开展了模型量化、"
        "FPGA算子设计、PS–PL数据流优化、Linux运行时管理与Web交互系统集成研究。项目以前期"
        "Qwen系列模型端侧部署经验为基础，进一步选取Microsoft Mage-VL 4B作为多模态目标模型，"
        "采用高精度CPU视觉编码与语言侧W4A8-G64量化加速相结合的混合执行方案。在硬件侧，针对"
        "849-token输入下Prefill矩阵计算量大的特点，提出M74六Token并行、六权重分片、两批"
        "反量化的算术候选，并在M73混合精确乘法结构上增加权重尺度预解码与精确移位加法路径。"
        "Vitis HLS 2026.1功能仿真覆盖W2、W3、W4三种位宽、6个Token、6个Shard与2个Batch，"
        "结果全部通过；核级估算周期为3.239 ns。按270 MHz及零额外开销计算，矩阵计算下界约为"
        "7640.645 ms，相对匹配CPU首Token基线10481.116 ms的理论加速比约为1.372倍。该结果"
        "属于功能仿真与理论预测，尚不能替代完整布线和实板端到端测量。项目最终形成了一套由模型"
        "语义门禁、混合精度量化、六端口DDR访问、M74 Prefill算术核、可回滚部署流程和网页交互"
        "组成的端侧多模态推理技术路线，为后续实板集成与性能优化提供了可复现基础。",
    )
    replace(
        doc.paragraphs[25],
        "Abstract: This project investigates multimodal large-model deployment on a resource-constrained "
        "embedded FPGA platform. It addresses model capacity, first-token computation, DDR bandwidth, and "
        "PS–PL coordination overhead through mixed-precision quantization, custom FPGA arithmetic, Linux "
        "runtime management, and web integration. Microsoft Mage-VL 4B is selected as the target model. "
        "High-precision CPU vision is combined with W4A8-G64 acceleration for the language path. The proposed "
        "M74 candidate processes six tokens over six weight shards and two dequantization batches, adding "
        "weight-scale predecoding and exact shift-add arithmetic to the previous hybrid design. Vitis HLS "
        "2026.1 C simulation passes W2, W3, and W4 cases for all configured tokens, shards, and batches. The "
        "estimated kernel period is 3.239 ns. At 270 MHz, the zero-overhead matrix-compute floor is 7,640.645 "
        "ms versus a matched CPU TTFT of 10,481.116 ms, corresponding to a theoretical 1.372x speedup. These "
        "results are simulation and analytical projections rather than board-level end-to-end measurements. "
        "The project establishes a reproducible route for later full-overlay implementation and matched board "
        "validation.",
    )
    replace(
        doc.paragraphs[27],
        "关键词：多模态大模型；Mage-VL；FPGA；PS–PL协同；W4A8；Prefill加速",
    )


def build_body(doc: Document) -> None:
    add_heading(doc, "1    绪论", 1)
    add_heading(doc, "1.1   研究背景", 2)
    add_p(
        doc,
        "近年来，大语言模型和视觉语言模型在文本生成、图像问答、视频理解与智能交互等任务中快速发展。"
        "随着模型由单一文本能力扩展到图像和视频输入，推理过程不仅要执行多层Transformer计算，还要"
        "完成视觉编码、跨模态特征组织和长序列Prefill，计算量与存储压力进一步增加。云端GPU能够"
        "提供较高吞吐率，但其成本、能耗、网络依赖和数据隐私问题限制了在教学终端、工业现场与边缘"
        "设备中的直接应用。",
    )
    add_p(
        doc,
        "端侧推理能够使模型与数据保留在本地，并为弱网或离线环境提供持续服务。但十亿参数级模型的"
        "权重、激活、KV Cache和运行时缓冲区会共同占用有限DDR；自回归生成又需要反复访问权重。"
        "对于视频模型，首Token前还需处理大量视觉Token，Prefill阶段往往成为决定交互体验的关键。"
        "若只把一个矩阵乘法算子移植到FPGA，而不处理量化精度、数据布局、DMA、缓存和系统调度，"
        "局部加速很难转化为端到端收益。",
    )
    add_p(
        doc,
        "AMD Kria KV260采用Zynq UltraScale+ MPSoC架构，片上集成64位ARM Cortex-A53处理系统"
        "和可编程逻辑。PS适合承担视频预处理、分词、模型控制、Attention、KV Cache与网页服务，"
        "PL则可以针对量化线性层建立高并行、低控制开销的数据通路。该平台资源与带宽均明显低于"
        "服务器FPGA，因而更适合研究“有限资源条件下如何获得真实端到端收益”这一工程问题。",
    )

    add_heading(doc, "1.2   国内外研究现状与问题分析", 2)
    add_p(
        doc,
        "现有大模型推理优化主要包括低比特量化、算子融合、KV Cache管理、权重驻留和异构调度。"
        "AWQ、SmoothQuant等方法通过降低权重或激活位宽减少存储与计算；vLLM通过PagedAttention"
        "改善服务器GPU的缓存管理与并发能力。此类方法为系统优化提供了重要思路，但服务器平台具有"
        "较大显存和成熟算子库，其并发服务目标与KV260上的单用户、低资源推理并不完全相同。",
    )
    add_p(
        doc,
        "在FPGA方向，FlightLLM、EdgeLLM等研究证明了定制矩阵阵列、混合精度和CPU–FPGA协同的"
        "可行性，但主要使用HBM或资源更丰富的高端器件。已有嵌入式FPGA工作往往针对更小模型、"
        "裸机系统或单个算子，较少同时给出模型语义、硬件调用、完整时序、稳定回滚和同精度CPU基线。"
        "因此，本项目将研究范围限定为：在KV260上构建可验证的Mage-VL 4B混合推理路径，并以"
        "首Token阶段的矩阵计算作为M74的主要优化对象。",
    )

    add_heading(doc, "1.3   研究目标与主要内容", 2)
    add_p(
        doc,
        "本项目总体目标是形成一套适用于KV260的多模态大模型软硬件协同方案，在保证模型语义和"
        "稳定回滚的前提下，验证FPGA相对于匹配CPU基线的性能潜力。主要内容包括：",
    )
    for line in [
        "1）建立Mage-VL 4B的模型结构、Token规模、计算量、存储容量和CPU首Token基线；",
        "2）完成语言侧W2/W3/W4混合低比特权重与A8-G64激活格式，保持视觉侧高精度执行；",
        "3）设计六端口DDR权重访问、语言权重高地址常驻及LM Head低地址映射的PS–PL架构；",
        "4）提出M74六Token Prefill算术候选，验证预解码移位加法的数值、资源和时序特征；",
        "5）建立包含制品哈希、Build ID、硬件计数、回答语义、性能数据和自动回滚的上板门禁；",
        "6）将本地推理服务接入Web前端，为后续视频问答和远程演示提供统一接口。",
    ]:
        add_p(doc, line)

    add_heading(doc, "1.4   技术路线与报告结构", 2)
    add_p(
        doc,
        "项目遵循“基线建立—量化校准—制品固化—算子仿真—综合实现—上板门禁—应用集成”的路线。"
        "第二章分析平台与模型约束；第三章介绍量化、内存布局和PS–PL划分；第四章重点说明M74结构"
        "及仿真；第五章介绍部署管理与Web交互；第六章总结成果、创新与风险；第七章给出结论和展望。",
    )

    add_heading(doc, "2    Mage-VL 4B与KV260平台分析", 1)
    add_heading(doc, "2.1   项目定位与核心技术问题", 2)
    add_p(
        doc,
        "项目定位不是在KV260上简单复制一个现成模型，而是建立从模型精度、存储布局到硬件执行证据"
        "的完整闭环。当前需要解决的核心问题包括：有限DDR如何容纳语言权重、视觉权重、LM Head与"
        "运行时缓冲；不同位宽的线性层如何在同一数据通路中统一处理；849-token视频输入如何在目标"
        "首Token时间内完成Prefill；以及局部HLS结果如何经过完整布线和实板门禁转化为可信结论。",
    )
    add_heading(doc, "2.2   软硬件协同总体架构", 2)
    add_p(
        doc,
        "系统分为应用交互层、PS运行时层、PL加速层和DDR存储层。应用层负责视频与提示词输入、状态"
        "显示和回答输出；PS层负责官方模型加载、视觉高精度计算、Token组织、Attention、残差、采样"
        "和硬件调度；PL层面向语言侧量化线性层和LM Head候选计算；存储层采用高低地址双池与六路"
        "128-bit权重主端口，减少单一AXI端口争用。",
    )
    add_table(
        doc,
        ["层次", "主要任务", "执行位置", "设计原则"],
        [
            ["应用层", "视频/提示词输入、网页展示", "主机与板端Web", "交互与推理解耦"],
            ["模型控制层", "视觉编码、Attention、采样", "Cortex-A53 PS", "高精度与灵活控制"],
            ["矩阵加速层", "语言线性层、LM Head候选", "KV260 PL", "W4A8与六分片并行"],
            ["存储层", "权重、激活、KV Cache、DMA", "DDR低/高地址池", "身份锁定与边界检查"],
        ],
    )
    add_heading(doc, "2.3   计算量与性能门禁", 2)
    add_p(
        doc,
        "锁定测试输入在首Token前包含约4.057×10^12次矩阵MAC和6.960×10^11次Attention MAC。"
        "匹配CPU TTFT为10481.116 ms；若要求至少1.25倍加速，则FPGA完整路径上限为8384.893 ms。"
        "因此仅复用200 MHz Decode核无法完成Prefill目标，必须采用接近全器件并行度、目标频率不低于"
        "约250 MHz的Token-tiled Prefill结构，并为Attention与运行时保留额外时间预算。",
    )

    add_heading(doc, "3    混合精度量化与六端口存储映射", 1)
    add_heading(doc, "3.1   量化策略与语义门禁", 2)
    add_p(
        doc,
        "项目对语言侧线性层采用W2/W3/W4混合位宽和A8-G64激活格式。每个64元素组先把激活尺度"
        "舍入为FP16，再生成有符号int8编码，并序列化为64字节编码、2字节尺度和14字节零填充。"
        "权重制品保存低比特编码、FP16尺度和矩阵方向，加载时逐文件核验尺寸与SHA-256，避免量化策略"
        "或存储方向在优化过程中被无意改变。",
    )
    add_p(
        doc,
        "全模型A8曾导致一个锁定视频语义错误，因此最终采用精度优先的混合路线：视觉编码器保持高精度"
        "CPU执行，252个语言线性层与untied LM Head进入A8路径。固定视频和三组方向/颜色语义门禁"
        "达到3/3通过。该结果证明所选量化作用域能够保持当前测试语义，但不等同于通用数据集准确率。",
    )
    add_heading(doc, "3.2   双池常驻与六端口访问", 2)
    add_p(
        doc,
        "语言权重按六个shard重排为group-major布局，并放入高地址保留DDR；LM Head单独形成六个"
        "低地址映射文件。语言侧页对齐映射约1694220288字节，LM Head映射约158048256字节。"
        "六个128-bit AXI主端口分别连接HPC0、HPC1、HP0、HP1、HP2和HP3，并同时配置低地址与"
        "高地址段。运行时在分配前验证manifest、文件哈希、物理地址窗口和tile边界，随后再编程PL。",
    )
    add_heading(doc, "3.3   正确性与生命周期控制", 2)
    add_p(
        doc,
        "候选部署采用fail-closed原则：任何文件哈希、地址范围、模型身份或输出检查失败，均在加载"
        "候选前停止；若候选已启动，则先恢复稳定boot，再安排重启并删除一次性服务。稳定方案和候选"
        "方案分别保存，避免把失败尝试覆盖为最新版本。由此形成“离线语义—HLS功能—Vivado实现—"
        "板端算子—整模型回答—性能配对—稳定回滚”的七级证据链。",
    )

    add_heading(doc, "4    M74 Prefill算术核设计与仿真", 1)
    add_heading(doc, "4.1   设计动机与前序方案分析", 2)
    add_p(
        doc,
        "M68首先验证六Token、六Shard并行算术路径，六个计算循环均达到II=1，但HLS估算为1296个"
        "DSP和141854个LUT，资源超限。M69降为五Token后DSP降至1080个，但理论加速上限不足以"
        "满足1.25倍目标。M70使用精确int-to-FP16转换减少DSP，却因转换网络完全复制导致LUT升至"
        "172156个。M71尝试三路两批共享，但HLS把结构补成六槽，资源反而恶化。",
    )
    add_p(
        doc,
        "M73回到六Token，并将三路反量化Token lane中的两路改为精确二进制16位fabric乘法，"
        "第三路保留原生half乘法。Vivado OOC综合显示83167个LUT、64108个FF、12个RAMB36和"
        "1248个DSP，资源可以放入孤立IP，但270 MHz时序为WNS=-0.113 ns，且DSP没有集成余量。"
        "M74因此进一步把可预处理的权重尺度结构化，尝试通过预解码与移位加法改善关键路径。",
    )
    add_heading(doc, "4.2   M74预解码移位加法结构", 2)
    add_p(
        doc,
        "M74仍保持六Token、六权重分片和两批反量化，避免降低矩阵吞吐。核心变化包括两部分：第一，"
        "在主循环外预解码FP16权重尺度，提取符号、指数与尾数相关信息，减少循环内重复判断；第二，"
        "对满足结构条件的乘法采用精确移位加法实现，并保持round-to-nearest-even语义。原生half"
        "路径仍被保留，用于覆盖不适合移位表达的情况。",
    )
    add_p(
        doc,
        "每个Shard内部完成A8编码与低比特权重的packed dot，随后把整数部分与预解码尺度组合。"
        "六路结果经共享反量化与输出阶段合并。该结构的目标并非单纯减少某一类算术单元，而是在不改变"
        "量化制品和功能输出的前提下，缩短反量化路径并为完整Prefill累加、Attention时分复用和系统"
        "集成保留空间。",
    )
    add_heading(doc, "4.3   功能仿真结果", 2)
    add_p(
        doc,
        "M74已在Vitis HLS 2026.1中完成C级仿真。测试覆盖W2、W3、W4三种权重方向与位宽，"
        "并同时覆盖6个Token、6个Shard和2个反量化Batch。日志返回MAGE_PREFILL_T6_"
        "SHARED_DEQUANT_CSIM_PASS，仿真错误数为0，流最大深度为36。由此可确认当前测试范围内"
        "的算术路径与软件参考保持一致。",
    )
    add_heading(doc, "4.4   HLS资源、时序与理论性能", 2)
    add_table(
        doc,
        ["指标", "M74当前结果", "结论边界"],
        [
            ["目标周期", "3.70 ns", "HLS约束"],
            ["估算周期", "3.239 ns（约308.7 MHz）", "核级估算，非布线频率"],
            ["BRAM18K", "48 / 288", "HLS估算"],
            ["DSP", "1248 / 1248", "已达器件上限"],
            ["FF", "172016 / 234240", "HLS估算"],
            ["LUT", "212020 / 117120", "HLS估算超限，须OOC复核"],
            ["功能", "W2/W3/W4 CSim PASS", "不是板端端到端结果"],
        ],
    )
    add_p(
        doc,
        "按270 MHz与零额外开销计算，六Token矩阵阵列的计算下界为7640.645 ms；相对CPU基线"
        "10481.116 ms，理论加速比为1.372倍，时间缩短约27.1%。该结果说明M74在矩阵阶段具有"
        "超过CPU 25%的理论空间，但距离完整1.25倍门禁只剩744.248 ms用于Attention、视觉预处理、"
        "数据搬运、PS–PL同步和运行时控制。因此必须把该数字标注为“理论下界”，不能写成实板实测。",
    )
    add_heading(doc, "4.5   当前状态与下一门禁", 2)
    add_p(
        doc,
        "M74当前状态为“功能仿真通过、OOC/完整系统与实板待验证”。下一步依次完成Vivado OOC综合、"
        "真实资源与270 MHz时序分析、完整输入组累加、Attention融合或时分复用、full-overlay布线、"
        "板端算子正确性以及与CPU同输入同精度的端到端配对测试。任何一级未通过均保留为否决记录，"
        "不进入下一阶段。",
    )

    add_heading(doc, "5    运行时管理与Web交互设计", 1)
    add_heading(doc, "5.1   Linux运行时与内存生命周期", 2)
    add_p(
        doc,
        "板端Linux运行时负责检查模型与硬件制品身份、申请连续内存、加载六分片权重、同步缓存、"
        "配置64位物理地址并启动PL。为降低文件页缓存和CMA竞争，权重按生命周期区分为长期常驻、"
        "逐请求激活和可释放临时数据；Prefill结束后只释放不再使用的页缓存，不回收Decode继续依赖的"
        "常驻权重。部署脚本要求候选失败后自动恢复稳定启动配置。",
    )
    add_heading(doc, "5.2   Web交互与运行证据展示", 2)
    add_p(
        doc,
        "Web后端采用轻量HTTP服务封装本地推理接口，模型在服务启动阶段一次性加载，并通过互斥锁保护"
        "共享KV Cache、CMA缓冲区和FPGA状态。前端使用HTML、CSS和JavaScript实现提示词输入、"
        "视频选择、生成长度、状态反馈、回答显示和错误提示。推理响应除文本外，还应返回模型身份、"
        "FPGA Build ID、PL调用计数、CPU回退状态和性能分解，使远程演示不仅“看见答案”，还能够"
        "确认硬件路径是否执行。",
    )
    add_heading(doc, "5.3   公网演示与安全边界", 2)
    add_p(
        doc,
        "系统可通过临时隧道提供公网访问，用于异地演示。该方式依赖板卡、主机、转发服务和隧道进程"
        "持续在线，不应视为长期生产部署。网页仅暴露受控推理接口，不提供Shell、文件写入或任意模型"
        "路径参数；并发请求返回忙状态，避免多个请求同时占用同一FPGA与内存池。",
    )
    add_heading(doc, "5.4   异地访问链路与演示流程", 2)
    add_p(
        doc,
        "为满足指导教师或项目成员异地访问的需求，网页连接采用“访问者浏览器—Cloudflare临时公网"
        "隧道—开发主机转发—KV260板端Web服务—本地模型运行时—FPGA加速核”的链路。板端只负责"
        "模型推理与状态返回，开发主机负责把局域网服务安全地转发到临时公网地址。访问者无需安装模型"
        "环境，只需打开链接、输入提示词或选择视频即可完成交互。",
    )
    add_table(
        doc,
        ["连接环节", "功能", "演示时检查内容"],
        [
            ["网页前端", "提示词/视频输入、回答与状态显示", "页面可打开、输入合法、截断有提示"],
            ["公网隧道", "生成异地可访问的HTTPS临时链接", "链接有效、无Shell与文件接口"],
            ["主机转发", "把公网请求转至板端局域网服务", "端口连通、超时和断线可提示"],
            ["板端Web后端", "健康检查、单请求调度、错误处理", "READY/BUSY状态正确"],
            ["模型与FPGA", "执行推理并返回硬件证据", "模型身份、Build ID、PL计数、无CPU回退"],
        ],
    )
    add_p(
        doc,
        "正式演示前应依次确认板端Linux与模型服务为READY、开发主机能够访问板端地址、隧道进程在线，"
        "再把当次生成的HTTPS链接发送给异地访问者。由于Quick Tunnel链接可能在重启后改变，报告中"
        "不固化某一个URL；若公网链路异常，则使用同一网页的局域网访问或预先录制的完整交互视频作为"
        "备用展示，但性能数据仍以板端原始日志为准。",
    )

    add_heading(doc, "6    阶段成果、创新点与风险分析", 1)
    add_heading(doc, "6.1   阶段性成果总结", 2)
    for line in [
        "1）建立Mage-VL 4B官方检查点、固定视频输入和10481.116 ms匹配CPU TTFT基线；",
        "2）形成W2/W3/W4混合量化制品和hash-verified fresh-process重载流程；",
        "3）确定视觉高精度、语言与LM Head A8的混合执行策略，锁定语义测试达到3/3；",
        "4）完成六端口LOW+HIGH DDR硬件映射及语言高地址、LM Head低地址的常驻布局；",
        "5）完成M68至M74多轮架构筛选，保留资源、时序和被否决方案的完整证据；",
        "6）M74在W2/W3/W4、6 Token、6 Shard和2 Batch条件下通过HLS功能仿真；",
        "7）建立从离线语义、综合实现到板端正确性、性能和回滚的分级验证体系；",
        "8）完成可承载本地模型的网页交互路线，为后续视频问答实板演示预留接口。",
    ]:
        add_p(doc, line)

    add_heading(doc, "6.2   项目创新点", 2)
    add_heading(doc, "6.2.1  面向端侧多模态模型的精度分域", 3)
    add_p(
        doc,
        "现有低比特部署常对全模型统一量化。本项目根据语义门禁结果，把视觉编码器保留在高精度CPU"
        "路径，只将语言线性层与LM Head映射到A8/低比特硬件路径，在有限内存下兼顾准确性与加速范围。",
    )
    add_heading(doc, "6.2.2  六端口双地址池与制品身份锁定", 3)
    add_p(
        doc,
        "项目将六路权重主端口同时映射到DDR低地址和高地址区域，配合语言权重与LM Head分池常驻。"
        "运行前必须通过manifest、SHA-256、地址范围和tile边界检查，从系统层面降低了错误权重、"
        "地址越界和候选覆盖稳定版本的风险。",
    )
    add_heading(doc, "6.2.3  M74预解码移位加法Prefill算术路径", 3)
    add_p(
        doc,
        "M74针对六Token Prefill反量化路径，在保持W2/W3/W4输出语义的前提下增加尺度预解码和"
        "精确移位加法，将量化格式、并行分片和时序约束联合设计。其价值不在于单纯堆叠并行度，而在于"
        "为完整累加与Attention融合寻找可实现的资源—时序平衡。",
    )
    add_heading(doc, "6.2.4  以否决记录和稳定回滚为核心的工程验证", 3)
    add_p(
        doc,
        "项目不把局部仿真、单算子速度或理论下界直接写成整模型性能。每个候选均记录Build ID、"
        "资源、时序、功能、语义、板端状态和下一门禁，并保留稳定回滚。这种证据链使优化过程能够被"
        "复核，也避免为了追求指标而掩盖失败尝试。",
    )

    add_heading(doc, "6.3   存在问题与风险预案", 2)
    add_table(
        doc,
        ["风险", "当前表现", "应对措施"],
        [
            ["资源不足", "M74 HLS估算DSP满载、LUT超限", "以OOC实值为准，降低转换复制并时分复用Attention"],
            ["270 MHz不收敛", "M73 OOC曾出现WNS=-0.113 ns", "预解码、流水寄存、层次约束并设较低频率退路"],
            ["端到端收益不足", "理论预算只剩744.248 ms", "优化DDR流、控制启动、Attention和PS–PL同步"],
            ["量化语义退化", "全模型A8曾有1例失败", "保留高精度视觉路径与3/3锁定语义门禁"],
            ["启动或候选故障", "候选可能无法进入稳定Linux", "稳定boot独立保存、一次性服务、串口与自动回滚"],
            ["公网演示不稳定", "临时隧道依赖多进程在线", "健康检查、超时提示、现场局域网/录屏备用方案"],
        ],
    )
    add_p(
        doc,
        "最重要的边界是：M74尚未完成full-overlay和实板端到端测试。因此本报告将7640.645 ms、"
        "1.372倍和27.1%统一标注为理论预测，不把它们表述为板端真实结果。实板结论必须同时包含"
        "运行日志、模型输出、FPGA执行证据和匹配CPU对照。",
    )

    add_heading(doc, "7    结论与展望", 1)
    add_heading(doc, "7.1   研究结论", 2)
    add_p(
        doc,
        "本项目围绕Mage-VL 4B在KV260上的端侧部署与首Token加速，完成了从模型语义、量化制品、"
        "六端口存储映射到M74算术候选的系统研究。混合精度路线通过锁定语义测试；六端口高低地址架构"
        "为语言权重与LM Head常驻提供了实现基础；M74在三种权重位宽、六Token、六Shard和两批"
        "反量化条件下通过HLS功能仿真，核级估算周期达到3.239 ns。",
    )
    add_p(
        doc,
        "性能分析显示，M74在270 MHz理想条件下具有1.372倍CPU加速潜力，但HLS资源估算和完整"
        "流水预算仍然紧张。由此得到的核心认识是：嵌入式FPGA大模型优化的本质不是单个乘法器更快，"
        "而是量化精度、并行阵列、DDR数据流、PS–PL职责和运行时生命周期的联合优化。当前成果已经"
        "形成可复现的设计与验证路线，但尚未形成M74实板性能结论。",
    )
    add_heading(doc, "7.2   后续工作", 2)
    for line in [
        "1）完成M74 Vivado OOC综合，确认真实LUT、DSP、BRAM、FF和270 MHz时序；",
        "2）增加完整输入组累加，并设计Attention融合或同DSP时分复用方案；",
        "3）完成full-overlay实现和板端算子正确性，禁止在此之前声称已部署；",
        "4）以相同检查点、输入、输出长度和计时边界完成CPU/FPGA重复配对测试；",
        "5）在性能通过后接入视频问答Web，并展示Build ID、调用计数和阶段耗时；",
        "6）将源码、量化制品、硬件工程、实验日志和失败候选整理为可验收成果。",
    ]:
        add_p(doc, line)
    add_p(
        doc,
        "若M74完整实现无法满足资源或1.25倍门禁，项目将保留当前结果作为算术结构研究结论，并回退"
        "到资源更可控的五Token路线，以“严格快于CPU且准确”为阶段目标，而不强行追求不可验证指标。",
    )

    add_p(doc, "参考文献", "Normal")
    references = [
        "[1] Vaswani A, Shazeer N, Parmar N, et al. Attention Is All You Need[C]. NeurIPS, 2017.",
        "[2] Frantar E, Ashkboos S, Hoefler T, et al. GPTQ: Accurate Post-Training Quantization for Generative Pre-trained Transformers[C]. ICLR, 2023.",
        "[3] Lin J, Tang J, Tang H, et al. AWQ: Activation-aware Weight Quantization for LLM Compression and Acceleration[C]. MLSys, 2024.",
        "[4] Xiao G, Lin J, Seznec M, et al. SmoothQuant: Accurate and Efficient Post-Training Quantization for Large Language Models[C]. ICML, 2023.",
        "[5] Kwon W, Li Z, Zhuang S, et al. Efficient Memory Management for Large Language Model Serving with PagedAttention[C]. SOSP, 2023.",
        "[6] Zeng S, Liu J, He G, et al. FlightLLM: Efficient Large Language Model Inference with a Complete Mapping Flow on FPGAs[C]. FPGA, 2024.",
        "[7] AMD. Kria KV260 Vision AI Starter Kit User Guide: UG1089[EB/OL].",
        "[8] AMD. Vitis High-Level Synthesis User Guide: UG1399[EB/OL].",
        "[9] AMD. Zynq UltraScale+ MPSoC Data Sheet: DS891[EB/OL].",
        "[10] Qwen Team. Qwen2.5 Technical Report[EB/OL]. 2024.",
    ]
    for ref in references:
        add_p(doc, ref)

    add_p(doc, "致  谢", "Normal")
    add_p(
        doc,
        "本项目的推进离不开指导教师、团队成员和实验平台的支持。感谢指导教师在项目选题、技术路线、"
        "硬件安全和阶段验收方面给予的指导；感谢团队成员在模型准备、量化验证、HLS设计、Vivado实现、"
        "板端运行与报告整理中的协作；感谢学校提供KV260开发板、计算环境和项目实践机会。项目中的每一"
        "项通过与否决记录都来自持续实验与复核，在此向所有提供帮助的老师和同学表示诚挚感谢。",
    )


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="m74_rewrite_") as tmp:
        clean = Path(tmp) / "clean_source.docx"
        sanitize_package(SOURCE, clean)
        doc = Document(clean)
        rewrite_front_matter(doc)
        keep_front_matter_only(doc)
        build_body(doc)
        doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
